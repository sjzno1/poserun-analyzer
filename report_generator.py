"""
PoseRun Analyzer - 报告生成器
基于Python-docx生成Word格式分析报告
支持关键帧标注截图、逐项评分、历史对比等高级功能
"""

import os
import json
import subprocess
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import numpy as np


class ReportGenerator:
    """跑姿分析报告生成器"""

    # 各技术项权重（合计 = 1.0），姿势跑法核心项（跨步/落地/膝盖）权重更高
    ITEM_WEIGHTS = {
        'foot_contact': 0.15,   # 落地 - 脚刚触地位置
        'knee_bend': 0.12,      # 膝关节 - 膝盖弯曲角度
        'stride': 0.20,         # 跨步 - 落地位置（姿势跑法核心）
        'symmetry': 0.10,       # 落地框架 - 膝盖重合度
        'shin_angle': 0.10,     # 小腿角度 - 地面夹角
        'push_off': 0.08,       # 离地 - 推蹬程度
        'leg_extension': 0.07,  # 伸腿 - 摆荡腿前伸
        'arm_swing': 0.04,      # 上身 - 摆臂
        'stability': 0.05,      # 晃动 - 身体稳定性
        'force_line': 0.05,     # 整体姿态 - 力线
        'elasticity': 0.02,     # 弹性 - 基础弹性
        'smoothness': 0.02,     # 流畅度 - 动作流畅
    }

    def __init__(self):
        self.doc = Document()

        # 设置默认字体
        style = self.doc.styles['Normal']
        style.font.name = 'Microsoft YaHei'
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(2)

        # 设置中文字体
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    # ------------------------------------------------------------------ #
    # 主入口
    # ------------------------------------------------------------------ #
    def generate_report(self, video_file: str, user_info: dict,
                        analysis_results: list, output_path: str = None,
                        key_frames_dir: str = None, fmt: str = 'docx') -> str:
        """
        生成完整的跑姿分析报告

        Args:
            video_file: 视频文件路径或名称
            user_info: 用户基本信息
            analysis_results: 姿态分析结果列表
            output_path: 输出文件路径（可选）
            key_frames_dir: 关键帧标注截图目录（可选，传入则嵌入报告）
            fmt: 报告格式 'docx' 或 'pdf'（PDF 用 reportlab 直接生成，
                 内置 STSong-Light 中文字体，无需 LibreOffice）

        Returns:
            生成的报告文件路径
        """
        fmt = (fmt or 'docx').lower()
        if fmt not in ('docx', 'pdf'):
            fmt = 'docx'

        # 计算综合指标
        metrics = self._calculate_comprehensive_metrics(analysis_results)
        self.metrics = metrics

        # 计算综合得分（含逐项明细）并写回 user_info（基本信息表使用）
        score_info = self._calculate_score(metrics)
        self.score_info = score_info
        user_info = dict(user_info)
        user_info['score'] = score_info['total']
        user_info['grade'] = score_info['grade']

        if output_path is None:
            # 默认写入 reports/ 目录（与 app.py 的 REPORTS_FOLDER 一致，
            # 用绝对路径避免受工作目录影响；容器内该目录挂载到 NAS）
            reports_dir = os.path.join(
                os.path.dirname(os.path.abspath(__file__)), 'reports')
            os.makedirs(reports_dir, exist_ok=True)
            name_part = self._safe_name(user_info.get('name') or '')
            ts = datetime.now().strftime('%Y%m%d_%H%M%S')
            prefix = f"报告_{name_part}_{ts}" if name_part else f"报告_{ts}"
            output_path = os.path.join(reports_dir, f"{prefix}.{fmt}")

        if fmt == 'pdf':
            return self._build_pdf(user_info, video_file, metrics,
                                   output_path, key_frames_dir)

        # 设置页面边距
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)

        # 1. 标题
        self._add_title(user_info)

        # 2. 基本信息表
        self._add_basic_info(user_info, video_file)

        # 3. 技术评分表（核心，含逐项得分）
        self._add_technics_table(metrics)

        # 4. 问题总结
        self._add_problem_summary(metrics)

        # 5. 改进建议
        self._add_improvement_suggestions(metrics)

        # 6. 关键帧标注截图（对应分析）
        if key_frames_dir and os.path.exists(key_frames_dir):
            self._add_keyframe_images(key_frames_dir)

        # 7. 历史对比（如果有历史记录）
        history_file = os.path.join(os.path.dirname(video_file) if video_file else '.', 'history.json')
        if os.path.exists(history_file):
            self._add_history_comparison(history_file)

        # 保存报告
        self.doc.save(output_path)
        return output_path

    # ------------------------------------------------------------------ #
    # PDF 报告（reportlab 直接生成，中文用内置 STSong-Light 字体）
    # ------------------------------------------------------------------ #
    @staticmethod
    def _safe_name(name: str) -> str:
        """把人名清洗成可安全用于文件名的字符串"""
        import re
        name = (name or '').strip()
        name = re.sub(r'[\\/:*?"<>|\s]+', '_', name)
        return name[:20]

    def _build_pdf(self, user_info: dict, video_file: str, metrics: dict,
                   output_path: str, key_frames_dir: str) -> str:
        """用 reportlab 生成 PDF 报告（与 docx 报告内容一致）"""
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_CENTER
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                        Table, TableStyle, Image as RLImage)

        # 内置中文字体（reportlab 自带 CID 字体，无需字体文件）
        if 'STSong-Light' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        FONT = 'STSong-Light'

        st_title = ParagraphStyle('t', fontName=FONT, fontSize=18, leading=24,
                                  alignment=TA_CENTER)
        st_sub = ParagraphStyle('s', fontName=FONT, fontSize=9, leading=12,
                                alignment=TA_CENTER,
                                textColor=colors.HexColor('#666666'))
        st_h1 = ParagraphStyle('h1', fontName=FONT, fontSize=13, leading=18,
                               spaceBefore=10, spaceAfter=4)
        st_body = ParagraphStyle('b', fontName=FONT, fontSize=10, leading=14)
        st_small = ParagraphStyle('sm', fontName=FONT, fontSize=8, leading=11,
                                  alignment=TA_CENTER,
                                  textColor=colors.HexColor('#666666'))
        st_prob = ParagraphStyle('p', fontName=FONT, fontSize=10, leading=15,
                                 textColor=colors.HexColor('#C00000'))
        st_sug = ParagraphStyle('g', fontName=FONT, fontSize=10, leading=15,
                                textColor=colors.HexColor('#008000'))

        doc = SimpleDocTemplate(
            output_path, pagesize=A4,
            topMargin=2 * cm, bottomMargin=2 * cm,
            leftMargin=2.2 * cm, rightMargin=2.2 * cm,
            title='姿势跑法技术评估报告',
            author=user_info.get('coach', '系统自动评估'))

        story = []
        # 1. 标题
        story.append(Paragraph('姿势跑法技术评估报告', st_title))
        story.append(Paragraph('POSE RUNNING TECHNIQUE ANALYSIS', st_sub))
        story.append(Spacer(1, 6))
        story.append(Paragraph('_' * 60, st_sub))
        story.append(Spacer(1, 8))

        # 2. 基本信息表
        basic_info = [
            ['评估项目', '姿势跑法技术评估'],
            ['学员姓名', str(user_info.get('name', '未知'))],
            ['评估教练', str(user_info.get('coach', '系统自动评估'))],
            ['综合得分', f"{user_info.get('score', 0)}分（{user_info.get('grade', '未评级')}）"],
            ['评估日期', datetime.now().strftime('%Y-%m-%d')],
            ['视频来源', os.path.basename(video_file) if video_file else '未提供'],
        ]
        tbl = Table(basic_info, colWidths=[4 * cm, 11.6 * cm])
        tbl.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), FONT),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F2F2F2')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 10))

        # 3. 技术评分表
        story.append(Paragraph('技术评分详情', st_h1))
        breakdown = metrics.get('breakdown', [])
        if breakdown:
            rows = [['环节', '分项', '评判结果', '得分', '建议']]
            for b in breakdown:
                rows.append([b['category'], b['item'], b['result'],
                             f"{int(b['score'])}/100", b.get('suggestion', '')])
            rows.append(['综合', '加权总分', metrics.get('grade', ''),
                         f"{metrics.get('total', 0)}/100", ''])
            t2 = Table(rows, colWidths=[1.9 * cm, 2.2 * cm, 3.6 * cm,
                                        1.8 * cm, 6.1 * cm])
            style = [
                ('FONTNAME', (0, 0), (-1, -1), FONT),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E8E8E8')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]
            for i, b in enumerate(breakdown, start=1):
                if b['score'] < 60:
                    style.append(('TEXTCOLOR', (3, i), (3, i),
                                  colors.HexColor('#C00000')))
                elif b['score'] < 80:
                    style.append(('TEXTCOLOR', (3, i), (3, i),
                                  colors.HexColor('#C87800')))
                else:
                    style.append(('TEXTCOLOR', (3, i), (3, i),
                                  colors.HexColor('#008000')))
            style.append(('FONTNAME', (0, len(rows) - 1), (-1, len(rows) - 1),
                          FONT))
            t2.setStyle(TableStyle(style))
            story.append(t2)
        else:
            story.append(Paragraph('（无有效分析数据）', st_body))
        story.append(Spacer(1, 8))

        # 4. 问题总结
        problems = metrics.get('problems', [])
        if problems:
            story.append(Paragraph('主要问题分析', st_h1))
            for p in problems:
                story.append(Paragraph(f'• {p}', st_prob))
            story.append(Spacer(1, 6))

        # 5. 改进建议
        suggestions = metrics.get('suggestions', [])
        if suggestions:
            story.append(Paragraph('改进建议', st_h1))
            for s in suggestions:
                story.append(Paragraph(f'✓ {s}', st_sug))
            story.append(Spacer(1, 6))

        # 6. 关键帧标注截图
        if key_frames_dir and os.path.exists(key_frames_dir):
            story.append(Paragraph('关键帧标注截图', st_h1))
            desc = ('以下为关键帧的姿态识别标注图：绿色线为骨骼连线，蓝色点为关键点，'
                    '右上角叠加了该帧的膝盖/小腿角度、跨步程度与跑步阶段，与上文评分一一对应。')
            story.append(Paragraph(desc, st_small))
            story.append(Spacer(1, 4))
            for img_file in sorted(Path(key_frames_dir).glob('keyframe_*.jpg')):
                try:
                    from PIL import Image as PILImage
                    with PILImage.open(img_file) as im:
                        iw, ih = im.size
                    max_w = 12 * cm
                    max_h = 16 * cm
                    scale = min(max_w / iw, max_h / ih, 1.0)
                    story.append(RLImage(str(img_file),
                                         width=iw * scale, height=ih * scale))

                    info_file = Path(key_frames_dir) / f"{img_file.stem}_info.json"
                    if info_file.exists():
                        with open(info_file, 'r', encoding='utf-8') as f:
                            info = json.load(f)
                        m = info.get('metrics', {})
                        phase_map = {'heel_strike': '脚跟触地', 'mid_stance': '承重中期',
                                     'toe_off': '蹬离前', 'swing': '摆动期'}
                        phase = phase_map.get(m.get('running_phase', ''),
                                              m.get('running_phase', ''))
                        caption = (f"帧 #{info.get('frame_index', info.get('frame_number', 0))}"
                                   f" · 时间 {info.get('timestamp', 'N/A')}"
                                   f" · 阶段 {phase}"
                                   f" · 膝盖 {m.get('knee_flexion', 0):.0f}°"
                                   f" · 小腿 {m.get('shin_angle', 0):.0f}°"
                                   f" · {m.get('stride_analysis', {}).get('stride_level', '')}")
                        story.append(Paragraph(caption, st_small))
                    story.append(Spacer(1, 6))
                except Exception as e:
                    print(f"PDF添加图片失败: {img_file.name}, 错误: {str(e)}")
                    continue

        # 7. 历史对比（与 docx 版同路径逻辑）
        history_file = os.path.join(
            os.path.dirname(video_file) if video_file else '.', 'history.json')
        if os.path.exists(history_file):
            try:
                with open(history_file, 'r', encoding='utf-8') as f:
                    history_data = json.load(f)
            except Exception:
                history_data = {}
            records = history_data.get('history', [])
            if len(records) >= 2:
                story.append(Paragraph('历史数据对比', st_h1))
                rows = [['评估日期', '膝盖角度', '小腿角度', '跨步情况', '综合得分']]
                for record in records[-5:]:
                    m = record.get('metrics', {})
                    rows.append([
                        record.get('date', m.get('date', 'N/A')),
                        f"{float(m.get('avg_knee', m.get('knee_angle', 0))) or 0:.1f}°",
                        f"{float(m.get('avg_shin', m.get('shin_angle', 0))) or 0:.1f}°",
                        m.get('stride_level', 'N/A'),
                        str(m.get('score', 'N/A')),
                    ])
                t3 = Table(rows, colWidths=[3.2 * cm, 2.8 * cm, 2.8 * cm,
                                            3.2 * cm, 3.6 * cm])
                t3.setStyle(TableStyle([
                    ('FONTNAME', (0, 0), (-1, -1), FONT),
                    ('FONTSIZE', (0, 0), (-1, -1), 9),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E8E8E8')),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('TOPPADDING', (0, 0), (-1, -1), 3),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ]))
                story.append(t3)

        doc.build(story)
        return output_path

    # ------------------------------------------------------------------ #
    # 文档各部分
    # ------------------------------------------------------------------ #
    def _add_title(self, user_info: dict):
        """添加报告标题"""
        title = self.doc.add_heading('姿势跑法技术评估报告', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle = self.doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run('POSE RUNNING TECHNIQUE ANALYSIS')
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 100, 100)

        self.doc.add_paragraph('_' * 70)

    def _add_basic_info(self, user_info: dict, video_file: str):
        """添加基本信息表格"""
        basic_info = [
            ['评估项目', '姿势跑法技术评估'],
            ['学员姓名', user_info.get('name', '未知')],
            ['评估教练', user_info.get('coach', '系统自动评估')],
            ['综合得分', f"{user_info.get('score', 0)}分（{user_info.get('grade', '未评级')}）"],
            ['评估日期', datetime.now().strftime('%Y-%m-%d')],
            ['视频来源', os.path.basename(video_file) if video_file else '未提供'],
        ]

        table = self.doc.add_table(rows=len(basic_info), cols=2, style='Table Grid')
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, row in enumerate(basic_info):
            for j, cell_text in enumerate(row):
                cell = table.rows[i].cells[j]
                cell.text = cell_text
                if j == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True

        self.doc.add_paragraph('')

    def _add_technics_table(self, metrics: dict):
        """添加技术评分表格（核心部分，含逐项得分与加权总分）"""
        self.doc.add_heading('技术评分详情', level=1)

        breakdown = metrics.get('breakdown', [])
        if not breakdown:
            self.doc.add_paragraph('（无有效分析数据）')
            return

        # 列：环节 | 分项 | 评判结果 | 得分 | 建议
        headers = ['环节', '分项', '评判结果', '得分', '建议']
        table = self.doc.add_table(rows=len(breakdown) + 2, cols=5, style='Table Grid')
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        for i, item in enumerate(breakdown):
            row = table.rows[i + 1]
            row.cells[0].text = item['category']
            row.cells[1].text = item['item']
            row.cells[2].text = item['result']
            row.cells[3].text = f"{int(item['score'])}/100"
            row.cells[4].text = item.get('suggestion', '')
            # 得分列着色：低分标红、中分标橙、高分标绿
            color_cell = row.cells[3]
            for paragraph in color_cell.paragraphs:
                for run in paragraph.runs:
                    if item['score'] < 60:
                        run.font.color.rgb = RGBColor(200, 0, 0)
                    elif item['score'] < 80:
                        run.font.color.rgb = RGBColor(200, 120, 0)
                    else:
                        run.font.color.rgb = RGBColor(0, 128, 0)
                    run.font.bold = True

        # 合计行
        total_row = table.rows[len(breakdown) + 1]
        total_row.cells[0].text = '综合'
        total_row.cells[1].text = '加权总分'
        total_row.cells[2].text = metrics.get('grade', '')
        total_row.cells[3].text = f"{metrics.get('total', 0)}/100"
        total_row.cells[4].text = ''
        for cell in total_row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        self.doc.add_paragraph('')

    def _add_problem_summary(self, metrics: dict):
        """添加问题总结"""
        problems = metrics.get('problems', [])
        if not problems:
            return

        self.doc.add_heading('主要问题分析', level=1)
        for problem in problems:
            p = self.doc.add_paragraph()
            run = p.add_run(f"• {problem}")
            run.font.color.rgb = RGBColor(200, 0, 0)

    def _add_improvement_suggestions(self, metrics: dict):
        """添加改进建议"""
        suggestions = metrics.get('suggestions', [])
        if not suggestions:
            return

        self.doc.add_heading('改进建议', level=1)
        for suggestion in suggestions:
            p = self.doc.add_paragraph()
            run = p.add_run(f"✓ {suggestion}")
            run.font.color.rgb = RGBColor(0, 128, 0)

    def _add_keyframe_images(self, key_frames_dir: str):
        """添加关键帧标注截图（对应各帧分析指标）"""
        self.doc.add_heading('关键帧标注截图', level=1)

        desc = self.doc.add_paragraph()
        run = desc.add_run(
            '以下为关键帧的姿态识别标注图：绿色线为骨骼连线，蓝色点为关键点，'
            '右上角叠加了该帧的膝盖/小腿角度、跨步程度与跑步阶段，与上文评分一一对应。')
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 100, 100)

        images_added = 0
        for img_file in sorted(Path(key_frames_dir).glob('keyframe_*.jpg')):
            try:
                self.doc.add_picture(str(img_file), width=Inches(4.0))
                images_added += 1

                info_file = Path(key_frames_dir) / f"{img_file.stem}_info.json"
                if info_file.exists():
                    with open(info_file, 'r', encoding='utf-8') as f:
                        info = json.load(f)
                    m = info.get('metrics', {})
                    phase_map = {'heel_strike': '脚跟触地', 'mid_stance': '承重中期',
                                 'toe_off': '蹬离前', 'swing': '摆动期'}
                    phase = phase_map.get(m.get('running_phase', ''), m.get('running_phase', ''))
                    caption = self.doc.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = caption.add_run(
                        f"帧 #{info.get('frame_index', info.get('frame_number', 0))} · 时间 {info.get('timestamp', 'N/A')}"
                        f" · 阶段 {phase}"
                        f" · 膝盖 {m.get('knee_flexion', 0):.0f}°"
                        f" · 小腿 {m.get('shin_angle', 0):.0f}°"
                        f" · {m.get('stride_analysis', {}).get('stride_level', '')}")
                    run.font.size = Pt(8)
                    run.font.italic = True
            except Exception as e:
                print(f"添加图片失败: {img_file.name}, 错误: {str(e)}")
                continue

        if images_added == 0:
            self.doc.add_paragraph('[关键帧截图区域：未提供标注图片]')

        self.doc.add_paragraph('')

    def _add_history_comparison(self, history_file: str):
        """添加历史数据对比"""
        try:
            with open(history_file, 'r', encoding='utf-8') as f:
                history_data = json.load(f)
        except Exception:
            return

        if not history_data or 'history' not in history_data:
            return

        records = history_data.get('history', [])
        if len(records) < 2:
            return

        self.doc.add_heading('历史数据对比', level=1)

        table = self.doc.add_table(rows=len(records[-5:]) + 1, cols=5, style='Table Grid')
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ['评估日期', '膝盖角度', '小腿角度', '跨步情况', '综合得分']
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        for i, record in enumerate(records[-5:]):
            row_idx = i + 1
            m = record.get('metrics', {})
            table.rows[row_idx].cells[0].text = record.get('date', m.get('date', 'N/A'))
            table.rows[row_idx].cells[1].text = f"{float(m.get('avg_knee', m.get('knee_angle', 0))) or 0:.1f}°"
            table.rows[row_idx].cells[2].text = f"{float(m.get('avg_shin', m.get('shin_angle', 0))) or 0:.1f}°"
            table.rows[row_idx].cells[3].text = m.get('stride_level', 'N/A')
            table.rows[row_idx].cells[4].text = str(m.get('score', 'N/A'))

        self.doc.add_paragraph('')

    # ------------------------------------------------------------------ #
    # 指标计算
    # ------------------------------------------------------------------ #
    def _calculate_comprehensive_metrics(self, analysis_results: list) -> dict:
        """计算综合指标（多帧统计 + 逐项评分明细）"""
        if not analysis_results:
            return {}

        knees = [r.get('knee_flexion', 0) for r in analysis_results]
        shins = [r.get('shin_angle', 0) for r in analysis_results]
        leangs = [r.get('body_leang', 0) for r in analysis_results]
        stride_dists = [r.get('stride_analysis', {}).get('stride_distance', 0)
                        for r in analysis_results]
        sym_scores = [r.get('symmetry', {}).get('symmetry_score', 50)
                      for r in analysis_results]
        hip_xs = [r.get('stride_analysis', {}).get('center_of_mass_x', 0.5)
                  for r in analysis_results]

        avg_knee = float(np.mean(knees))
        avg_shin = float(np.mean(shins))
        avg_leang = float(np.mean(leangs))
        avg_stride = float(np.median(stride_dists))   # 用中位数，抗单帧异常
        avg_sym = float(np.mean(sym_scores))
        hip_std = float(np.std(hip_xs))              # 重心左右晃动幅度

        stride_level = self._classify_stride(avg_stride)
        symmetry_level = self._classify_symmetry(avg_sym)

        stride_analysis = {
            'stride_distance': avg_stride,
            'stride_level': stride_level,
        }
        symmetry = {
            'symmetry_score': avg_sym,
            'symmetry_level': symmetry_level,
        }

        # 逐项评分明细
        breakdown = self._build_breakdown(
            avg_knee, avg_shin, avg_leang, avg_stride, avg_sym, hip_std)

        # 加权总分
        total = self._weighted_total(breakdown)
        grade = self._grade_of(total)

        problems = [b['result'] for b in breakdown if b['score'] < 60]
        suggestions = [b['suggestion'] for b in breakdown
                       if b['score'] < 80 and b.get('suggestion')]

        return {
            'avg_knee_flexion': avg_knee,
            'avg_shin_angle': avg_shin,
            'avg_body_leang': avg_leang,
            'stride_analysis': stride_analysis,
            'symmetry': symmetry,
            'hip_stability': hip_std,
            'breakdown': breakdown,
            'total': total,
            'grade': grade,
            'problems': problems,
            'suggestions': suggestions,
        }

    # ------------------------------------------------------------------ #
    # 逐项评分
    # ------------------------------------------------------------------ #
    def _interp(self, value: float, points: list) -> float:
        """分段线性插值。points 为按 x 升序的 [(x, y), ...]"""
        if not points:
            return 0.0
        if value <= points[0][0]:
            return float(points[0][1])
        if value >= points[-1][0]:
            return float(points[-1][1])
        for (x0, y0), (x1, y1) in zip(points, points[1:]):
            if x0 <= value <= x1:
                t = (value - x0) / (x1 - x0) if x1 > x0 else 0.0
                return float(y0 + t * (y1 - y0))
        return float(points[-1][1])

    def _classify_stride(self, dist: float) -> str:
        if dist < 0.05:
            return '轻度跨步'
        if dist < 0.10:
            return '中度跨步'
        return '严重跨步'

    def _classify_symmetry(self, score: float) -> str:
        if score >= 80:
            return '良好'
        if score >= 60:
            return '一般'
        return '较差'

    def _build_breakdown(self, avg_knee, avg_shin, avg_leang,
                         avg_stride, avg_sym, hip_std) -> list:
        """对 12 个技术项逐项打分（0-100，越高越好）"""
        items = []

        # 1. 落地 - 脚刚触地位置（前掌/中足触地为佳，小腿越竖直越好）
        fc_score = self._interp(avg_shin, [(60, 30), (70, 55), (80, 80), (90, 100)])
        if avg_shin < 60:
            fc_result, fc_sug = '明显足跟触地', '落地点前移，让脚落在重心正下方'
        elif avg_shin < 75:
            fc_result, fc_sug = '轻度脚跟触地', '缩短跨步，前掌先着地'
        else:
            fc_result, fc_sug = '前掌/中足触地', ''
        items.append(self._mk('foot_contact', '落地', '脚刚触地位置',
                              fc_result, fc_score, fc_sug))

        # 2. 膝关节 - 膝盖弯曲角度（触地时宜微弯，既不过直也不过度屈）
        kb_score = self._interp(avg_knee, [(130, 30), (145, 60), (155, 80),
                                           (165, 100), (175, 100), (180, 70)])
        if avg_knee < 150:
            kb_result, kb_sug = '膝盖弯曲角度较大', '触地时膝盖放松微弯即可，避免过度屈膝缓冲'
        elif avg_knee < 165:
            kb_result, kb_sug = '膝盖有微弯（良好）', ''
        elif avg_knee <= 175:
            kb_result, kb_sug = '膝盖微弯适中', ''
        else:
            kb_result, kb_sug = '膝盖偏直（偏僵硬）', '落地时膝盖保持放松微弯'
        items.append(self._mk('knee_bend', '膝关节', '膝盖弯曲角度',
                              kb_result, kb_score, kb_sug))

        # 3. 跨步 - 落地位置（姿势跑法核心，落脚点越靠近重心越好）
        st_score = self._interp(avg_stride, [(0.02, 100), (0.05, 85),
                                             (0.10, 55), (0.15, 35), (0.20, 20)])
        st_level = self._classify_stride(avg_stride)
        st_result = f'跨步程度：{st_level}'
        st_sug = '收紧身体框架，脚落在重心正下方' if avg_stride >= 0.05 else ''
        items.append(self._mk('stride', '跨步', '落地位置',
                              st_result, st_score, st_sug))

        # 4. 落地框架 - 膝盖重合度（左右对称）
        sy_level = self._classify_symmetry(avg_sym)
        sy_result = f'对称性：{sy_level}'
        sy_sug = '注意左右腿对称发力' if avg_sym < 60 else ''
        items.append(self._mk('symmetry', '落地框架', '膝盖重合度',
                              sy_result, avg_sym, sy_sug))

        # 5. 小腿角度 - 地面夹角
        sh_score = self._interp(avg_shin, [(60, 35), (70, 55), (80, 80), (90, 100)])
        sh_result = f"{avg_shin:.1f}°"
        sh_sug = '小腿着地时保持竖直' if avg_shin < 70 else ''
        items.append(self._mk('shin_angle', '小腿角度', '地面夹角',
                              sh_result, sh_score, sh_sug))

        # 6. 离地 - 推蹬程度（姿势跑法主张少推蹬，靠重力前移）
        po_score = self._interp(avg_shin, [(60, 45), (65, 55), (75, 85), (85, 95)])
        po_result = '无明显推蹬' if avg_shin >= 70 else '存在推蹬'
        po_sug = '减少主动蹬地，靠重力带动身体前移' if avg_shin < 70 else ''
        items.append(self._mk('push_off', '离地', '推蹬程度',
                              po_result, po_score, po_sug))

        # 7. 伸腿 - 摆荡腿前伸（适度前倾 5-10° 最佳）
        le_score = self._interp(avg_leang, [(0, 45), (5, 90), (8, 100),
                                            (12, 85), (20, 55)])
        if avg_leang < 3:
            le_result, le_sug = '前倾不足，重力利用不够', '身体适度前倾，让重力帮你前进'
        elif avg_leang <= 12:
            le_result, le_sug = '前倾时机适中', ''
        else:
            le_result, le_sug = '前倾过度', '减小前倾角度，保持躯干稳定'
        items.append(self._mk('leg_extension', '伸腿', '摆荡腿前伸',
                              le_result, le_score, le_sug))

        # 8. 上身 - 摆臂（当前数据有限，给基准分）
        items.append(self._mk('arm_swing', '上身', '摆臂',
                              '基本正常', 78, ''))

        # 9. 晃动 - 身体稳定性（重心左右晃动越小越稳）
        sb_score = self._interp(hip_std, [(0.01, 100), (0.03, 75),
                                          (0.06, 50), (0.10, 25)])
        sb_result = '无明显晃动' if hip_std < 0.05 else '存在左右晃动'
        sb_sug = '收紧核心，减少躯干左右摆动' if hip_std >= 0.05 else ''
        items.append(self._mk('stability', '晃动', '身体稳定性',
                              sb_result, sb_score, sb_sug))

        # 10. 整体姿态 - 力线（前倾 + 对称综合）
        fl_score = 0.5 * le_score + 0.5 * avg_sym
        fl_result = '正常' if fl_score >= 80 else '待优化'
        fl_sug = '保持躯干中立、力线垂直' if fl_score < 80 else ''
        items.append(self._mk('force_line', '整体姿态', '力线',
                              fl_result, fl_score, fl_sug))

        # 11. 弹性 - 基础弹性（基准分）
        items.append(self._mk('elasticity', '弹性', '基础弹性',
                              '基本正常', 72, ''))

        # 12. 流畅度 - 动作流畅（基准分）
        items.append(self._mk('smoothness', '流畅度', '动作流畅',
                              '基本流畅', 74, ''))

        return items

    def _mk(self, key, category, item, result, score, suggestion) -> dict:
        return {
            'key': key,
            'category': category,
            'item': item,
            'result': result,
            'score': max(0.0, min(100.0, float(score))),
            'weight': self.ITEM_WEIGHTS.get(key, 0.0),
            'suggestion': suggestion,
        }

    def _weighted_total(self, breakdown: list) -> int:
        """加权总分（权重合计为1）"""
        w_sum = sum(b['weight'] for b in breakdown)
        if w_sum <= 0:
            return 0
        total = sum(b['score'] * b['weight'] for b in breakdown) / w_sum
        return int(round(max(0.0, min(100.0, total))))

    def _grade_of(self, total: int) -> str:
        if total >= 90:
            return '优秀'
        if total >= 80:
            return '良好'
        if total >= 70:
            return '一般'
        if total >= 55:
            return '待改进'
        return '较差'

    def _calculate_score(self, metrics: dict) -> dict:
        """计算综合得分，返回 {total, grade, breakdown}

        兼容旧调用：若需要纯整数，请使用 result['total']。
        """
        breakdown = metrics.get('breakdown')
        if not breakdown:
            breakdown = self._build_breakdown(
                metrics.get('avg_knee_flexion', 0),
                metrics.get('avg_shin_angle', 0),
                metrics.get('avg_body_leang', 0),
                metrics.get('stride_analysis', {}).get('stride_distance', 0),
                metrics.get('symmetry', {}).get('symmetry_score', 50),
                metrics.get('hip_stability', 0.0),
            )
        total = self._weighted_total(breakdown)
        return {
            'total': total,
            'grade': self._grade_of(total),
            'breakdown': breakdown,
        }

    def export_to_pdf(self, docx_path: str, pdf_path: str = None) -> str:
        """导出为PDF格式"""
        if pdf_path is None:
            pdf_path = os.path.splitext(docx_path)[0] + '.pdf'

        try:
            result = subprocess.run([
                'libreoffice', '--headless', '--convert-to', 'pdf',
                '--outdir', os.path.dirname(pdf_path) or '.',
                docx_path
            ], capture_output=True, text=True, timeout=60)

            if result.returncode == 0:
                return pdf_path
            else:
                print(f"PDF转换失败: {result.stderr}")
                return None
        except Exception as e:
            print(f"PDF转换错误: {str(e)}")
            return None


if __name__ == '__main__':
    # 测试报告生成
    generator = ReportGenerator()

    user_info = {
        'name': '测试用户',
        'coach': '系统自动评估',
    }

    sample_results = [
        {
            'knee_flexion': 158.5,
            'shin_angle': 74.2,
            'body_leang': 7.0,
            'stride_analysis': {'stride_distance': 0.06, 'stride_level': '中度跨步'},
            'symmetry': {'symmetry_score': 78, 'symmetry_level': '良好'},
        } for _ in range(10)
    ]

    output_file = generator.generate_report(
        'test_video.mp4',
        user_info,
        sample_results,
        '测试报告.docx'
    )

    print(f"报告已生成: {output_file}")
    print(f"综合得分: {generator.metrics['total']}（{generator.metrics['grade']}）")
